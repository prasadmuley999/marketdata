import os
import re
import glob
import pandas as pd
import openpyxl
import yfinance as yf
from datetime import datetime, timedelta
from openpyxl.utils import get_column_letter

# --- UPDATED: Dynamic dated filename ---
# Generates a string like "NSE_Merged_Reports_2026-09-09.xlsx" 
CURRENT_DATE_STR = datetime.now().strftime('%Y-%m-%d')
EXCEL_FILE = f"NSE_Merged_Reports_{CURRENT_DATE_STR}.xlsx"

# The specific list of 54 target symbols provided
TARGET_SYMBOLS = [
    "ADANIENT", "ADANIPORTS", "APOLLOHOSP", "ASIANPAINT", "AXISBANK", 
    "BAJAJ-AUTO", "BAJFINANCE", "BAJAJFINSV", "BEL", "BPCL", 
    "BHARTIARTL", "BRITANNIA", "CIPLA", "COALINDIA", "DIVISLAB", 
    "DRREDDY", "EICHERMOT", "ETERNAL", "GRASIM", "HCLTECH", 
    "HDFCBANK", "HDFCLIFE", "HEROMOTOCO", "HINDALCO", "HINDUNILVR", 
    "ICICIBANK", "INDIGO", "INDUSINDBK", "INFY", "JIOFIN", 
    "JSWSTEEL", "KOTAKBANK", "LT", "M&M", "MARUTI", 
    "MAXHEALTH", "NESTLEIND", "NTPC", "ONGC", "POWERGRID", 
    "RELIANCE", "SBILIFE", "SBIN", "SHRIRAMFIN", "SUNPHARMA", 
    "TATACONSUM", "TATAMOTORS", "TATASTEEL", "TCS", "TECHM", 
    "TITAN", "TRENT", "ULTRACEMCO", "WIPRO"
]

def parse_sheet_date(sheet_name):
    """Parses tab date format 'DD-MMM-YYYY' to sort sheets chronologically."""
    match = re.match(r'(\d{2})-([A-Za-z]{3})-(\d{4})', sheet_name)
    if match:
        try:
            return datetime.strptime(sheet_name, '%d-%b-%Y')
        except ValueError:
            pass
    return datetime.min

def fetch_historical_price_changes(dates_list):
    """Fetches end-of-day price changes from Yahoo Finance for all symbols over the target date range."""
    if not dates_list:
        return pd.DataFrame()
        
    min_date = min(dates_list)
    max_date = max(dates_list)
    
    # Pad range slightly to ensure we capture the prior close price for the oldest date
    start_date_str = (min_date - timedelta(days=7)).strftime('%Y-%m-%d')
    end_date_str = (max_date + timedelta(days=3)).strftime('%Y-%m-%d')
    
    tickers = [f"{sym}.NS" for sym in TARGET_SYMBOLS]
    print(f"Downloading historical pricing from {start_date_str} to {end_date_str}...")
    
    try:
        price_data = yf.download(tickers, start=start_date_str, end=end_date_str, progress=False)
        close_df = price_data['Close'] if 'Close' in price_data else price_data
        
        # Calculate daily percentage change
        change_df = close_df.pct_change() * 100
        
        # Format index to date-only to avoid timezone mismatch errors
        change_df.index = pd.to_datetime(change_df.index).date
        return change_df
    except Exception as e:
        print(f"Error fetching historical prices: {e}")
        return pd.DataFrame()

def generate_dashboard():
    # NOTE: If your 'export_to_excel.py' still outputs the old filename, 
    # you may need to look for "NSE_Merged_Reports.xlsx" instead.
    # Below assumes you want to look for the newly named file:
    if not os.path.exists(EXCEL_FILE):
        # Fallback check if export_to_excel.py hasn't been updated yet:
        OLD_FILE = "NSE_Merged_Reports.xlsx"
        if os.path.exists(OLD_FILE):
            print(f"Found base file '{OLD_FILE}', renaming to target '{EXCEL_FILE}'...")
            os.rename(OLD_FILE, EXCEL_FILE)
        else:
            print(f"Error: '{EXCEL_FILE}' or '{OLD_FILE}' not found. Please run 'export_to_excel.py' first.")
            return

    print(f"Loading '{EXCEL_FILE}'...")
    wb = openpyxl.load_workbook(EXCEL_FILE)

    # Remove any existing Dashboard sheet to avoid duplicates
    if "Dashboard" in wb.sheetnames:
        del wb["Dashboard"]

    # Gather and sort all daily data sheets (Newest First)
    daily_sheets = [s for s in wb.sheetnames if s != "Dashboard"]
    daily_sheets.sort(key=parse_sheet_date, reverse=True)

    if not daily_sheets:
        print("No daily worksheets found in the workbook.")
        return

    # Parse dates to fetch the matching historical close price data
    valid_dates = [parse_sheet_date(s) for s in daily_sheets if parse_sheet_date(s) != datetime.min]
    change_df = fetch_historical_price_changes(valid_dates)

    print(f"Generating Dashboard with {len(daily_sheets)} daily sheets...")

    # Create the Dashboard as the very first sheet (index 0)
    ws = wb.create_sheet(title="Dashboard", index=0)

    # --- UPDATED: Freeze the first column (Column A) ---
    # Freezing 'B2' keeps Column A locked for horizontal scrolling, 
    # and keeps Row 1 (headers) locked for vertical scrolling.
    # If you ONLY want to freeze the column and not the header row, use 'B1'.
    ws.freeze_panes = 'B2'

    # 1. Write headers into Row 1
    ws.cell(row=1, column=1, value="SYMBOL")
    
    for i, sheet_name in enumerate(daily_sheets):
        col_qty = 2 + (3 * i)  # Columns B, E, H, etc.
        col_pct = 3 + (3 * i)  # Columns C, F, I, etc.
        col_prc = 4 + (3 * i)  # Columns D, G, J, etc.
        
        ws.cell(row=1, column=col_qty, value=f"Deliv Qty ({sheet_name})")
        ws.cell(row=1, column=col_pct, value=f"Deliv % ({sheet_name})")
        ws.cell(row=1, column=col_prc, value=f"Price Change % ({sheet_name})")

    # 2. Populate symbols, formulas, and price values
    for r_idx, symbol in enumerate(TARGET_SYMBOLS, start=2):
        # Write the active symbol in Column A
        ws.cell(row=r_idx, column=1, value=symbol)
        
        # Populate columns across the daily intervals
        for i, sheet_name in enumerate(daily_sheets):
            col_qty = 2 + (3 * i)
            col_pct = 3 + (3 * i)
            col_prc = 4 + (3 * i)
            
            # VLOOKUP formulas pointing back to original daily sheets
            formula_qty = f"=IFERROR(VLOOKUP(A{r_idx}, '{sheet_name}'!A:E, 4, FALSE), \"-\")"
            formula_pct = f"=IFERROR(VLOOKUP(A{r_idx}, '{sheet_name}'!A:E, 5, FALSE), \"-\")"
            
            ws.cell(row=r_idx, column=col_qty, value=formula_qty)
            ws.cell(row=r_idx, column=col_pct, value=formula_pct)
            
            # Retrieve corresponding date's price change value
            dt = parse_sheet_date(sheet_name)
            target_date = dt.date()
            ticker = f"{symbol}.NS"
            
            price_change = None
            if not change_df.empty and ticker in change_df.columns:
                if target_date in change_df.index:
                    val = change_df.loc[target_date, ticker]
                    # Handle Series output safely if duplicates occur
                    if isinstance(val, pd.Series):
                        val = val.iloc[0]
                    if pd.notna(val):
                        price_change = val

            # Write price change value
            if price_change is not None:
                # Divide by 100 and apply Excel standard percentage format
                # (This keeps the value numeric so you can calculate/chart it)
                cell_value = price_change / 100.0
                cell = ws.cell(row=r_idx, column=col_prc, value=cell_value)
                cell.number_format = '+0.00%;-0.00%;0.00%'
            else:
                ws.cell(row=r_idx, column=col_prc, value="-")

    # 3. Adjust column widths automatically for clean formatting
    print("Auto-fitting column sizes...")
    for col in ws.columns:
        max_len = max(len(str(cell.value or '')) for cell in col)
        col_letter = get_column_letter(col[0].column)
        ws.column_dimensions[col_letter].width = max(max_len + 3, 12)

    # 4. Save workbook
    wb.save(EXCEL_FILE)
    print(f"\nSuccess: 'Dashboard' sheet compiled as Sheet1 in '{EXCEL_FILE}' with frozen panels!")

if __name__ == '__main__':
    generate_dashboard()
